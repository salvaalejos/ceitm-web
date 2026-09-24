from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from sqlmodel import Session, select, SQLModel, func
from sqlalchemy.orm import selectinload
from datetime import datetime, timedelta
from io import BytesIO
import re

from openpyxl import load_workbook

from app.core.database import get_session
from app.models.user_model import User, UserRole, UserArea
from app.models.student_model import Student
from app.models.career_model import Career
from app.models.scholarship_model import ScholarshipApplication, ApplicationStatus
from app.models.attendance_model import Attendance, AttendanceStatus
from app.models.servicio_becario_model import ServicioBecario, ServicioPeriodo, ServicioEstado
from app.api.deps import get_current_user

router = APIRouter()


# --- FUNCIÓN AUXILIAR DE PERMISOS ---
def is_becarios_manager(user: User) -> bool:
    """
    Permite el acceso al padrón de becarios a:
    Admin, Mesa Directiva, Presidencia, Área de Becas y Prevención y Logística.
    """
    if user.role in [UserRole.ADMIN_SYS, UserRole.ESTRUCTURA]:
        return True
    if user.area in [UserArea.BECAS, UserArea.PREVENCION, UserArea.PRESIDENCIA]:
        return True
    return False


def _check_becarios_manager(user: User):
    if not is_becarios_manager(user):
        raise HTTPException(status_code=403, detail="No tienes autorización para gestionar el padrón de becarios.")


# --- ESQUEMAS DE RESPUESTA LIGEROS ---
class CareerRelRead(SQLModel):
    name: str


class StudentReadWithCareer(SQLModel):
    control_number: str
    full_name: str
    email: str
    career: Optional[str] = None
    is_blacklisted: bool
    career_rel: Optional[CareerRelRead] = None
    scholarship_type: Optional[str] = None
    days_active: Optional[int] = None
    total_services: Optional[int] = 0
    released_services: Optional[int] = 0
    current_week_faults: Optional[int] = 0


class PaginatedStudents(SQLModel):
    total: int
    items: List[StudentReadWithCareer]


# --- ESQUEMAS DE SERVICIOS ---
class ServicioRead(SQLModel):
    id: int
    control_number: str
    actividad: str
    periodo: str
    anio: int
    estado: str
    folio: Optional[str] = None
    liberado_at: Optional[datetime] = None
    created_at: datetime


class ServicioCreate(SQLModel):
    actividad: str
    periodo: ServicioPeriodo
    anio: int
    liberar: bool = True  # Si True, se genera folio y pasa a LIBERADO


class ServicioUpdate(SQLModel):
    # Permite revertir una liberación (folios a None) o registrar como pendiente
    liberar: Optional[bool] = None
    actividad: Optional[str] = None
    periodo: Optional[ServicioPeriodo] = None
    anio: Optional[int] = None


# --- ESQUEMAS PARA ALTA MANUAL / IMPORTACIÓN ---
class StudentManualCreate(SQLModel):
    control_number: str
    full_name: str
    email: Optional[str] = None
    phone_number: Optional[str] = None
    career: str  # Nombre de la carrera (obligatoria)


class ImportExcelResponse(SQLModel):
    total: int
    creados: int
    actualizados: int
    duplicados: int
    servicios: int
    errores: List[str]


def _to_servicio_read(s: ServicioBecario) -> ServicioRead:
    return ServicioRead(
        id=s.id,
        control_number=s.control_number,
        actividad=s.actividad,
        periodo=s.periodo.value if hasattr(s.periodo, "value") else str(s.periodo),
        anio=s.anio,
        estado=s.estado.value if hasattr(s.estado, "value") else str(s.estado),
        folio=s.folio,
        liberado_at=s.liberado_at,
        created_at=s.created_at,
    )


def _normalize_text(value: Optional[str]) -> str:
    """Minúsculas, sin acentos y con espacios colapsados."""
    text = (value or "").strip().lower()
    for accented, plain in (("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"), ("ü", "u")):
        text = text.replace(accented, plain)
    return " ".join(text.split())


def _resolve_career_id(session: Session, name: Optional[str]) -> Optional[int]:
    """Busca la carrera por nombre sin distinguir mayúsculas, acentos ni espacios."""
    if not name or not name.strip():
        return None
    target = _normalize_text(name)
    careers = session.exec(select(Career)).all()
    for career in careers:
        if _normalize_text(career.name) == target:
            return career.id
    # Coincidencia parcial segura: solo si hay un único candidato
    candidates = [
        c for c in careers
        if target in _normalize_text(c.name) or _normalize_text(c.name) in target
    ]
    if len(candidates) == 1:
        return candidates[0].id
    return None


def _find_student(session: Session, control: str) -> Optional[Student]:
    """Busca al alumno por número de control ignorando mayúsculas/minúsculas."""
    student = session.get(Student, control)
    if student:
        return student
    return session.exec(
        select(Student).where(func.lower(Student.control_number) == control.lower())
    ).first()


def _generate_folio(servicio: ServicioBecario) -> str:
    return f"CEITM-{servicio.anio}{servicio.periodo.value}-{servicio.id:05d}"


def _derive_periodo_anio(folio: Optional[str], now: datetime):
    """Deduce periodo y año del folio (CEITM-AAAA[ABV]-...) o usa el mes actual."""
    if folio:
        match = re.match(r"CEITM-(\d{4})([ABV])", folio.strip().upper())
        if match:
            return ServicioPeriodo(match.group(2)), int(match.group(1))
    if now.month <= 6:
        return ServicioPeriodo.A, now.year
    if now.month == 7:
        return ServicioPeriodo.V, now.year
    return ServicioPeriodo.B, now.year


# --- ENDPOINTS ---
@router.get("/", response_model=PaginatedStudents)
def read_students(
        skip: int = Query(0, ge=0),
        limit: int = Query(10, ge=1, le=100),
        search: Optional[str] = Query(None),
        sort_by: Optional[str] = Query("control_desc"),
        session: Session = Depends(get_session),
        current_user: User = Depends(get_current_user),
):
    try:
        # 👇 1. Validación de Permisos Segura
        if not is_becarios_manager(current_user):
            raise HTTPException(status_code=403, detail="No tienes autorización para ver el padrón de becarios.")

        # 👇 NUEVO: Fechas para calcular la semana actual
        today = datetime.utcnow().date()
        start_of_week = today - timedelta(days=today.weekday())
        end_of_week = start_of_week + timedelta(days=4)

        # 👇 2. Consulta Base: TODOS los estudiantes (manuales, importados o con solicitud)
        base_query = select(Student).distinct()

        # 👇 3. Búsqueda
        if search:
            base_query = base_query.where(
                (Student.full_name.icontains(search)) |
                (Student.control_number.icontains(search))
            )

        # 👇 4. Ordenamiento Dinámico
        if sort_by == "name_asc":
            base_query = base_query.order_by(Student.full_name.asc())
        elif sort_by == "name_desc":
            base_query = base_query.order_by(Student.full_name.desc())
        elif sort_by == "control_asc":
            base_query = base_query.order_by(Student.control_number.asc())
        else:
            base_query = base_query.order_by(Student.control_number.desc())

        all_matching_students = session.exec(base_query).all()
        total = len(all_matching_students)

        query = base_query.options(
            selectinload(Student.career_rel),
            selectinload(Student.applications).selectinload(ScholarshipApplication.scholarship),
            selectinload(Student.servicios),
        ).offset(skip).limit(limit)

        students_db = session.exec(query).all()

        items = []
        now = datetime.utcnow()

        for student in students_db:
            active_app = None
            total_serv = 0
            released_serv = 0

            # Servicios (ServicioBecario) - totales y liberados
            for serv in student.servicios:
                total_serv += 1
                if serv.estado == ServicioEstado.LIBERADO:
                    released_serv += 1

            # Solicitudes (de lectura, para tipo de beca / días activos)
            for app in student.applications:
                if app.status in [ApplicationStatus.APROBADA, ApplicationStatus.LIBERADA]:
                    if not active_app or (
                            app.created_at and active_app.created_at and app.created_at > active_app.created_at):
                        active_app = app

            s_type = None
            d_active = None

            if active_app and active_app.scholarship:
                # Extracción súper segura del ENUM (Evita crash 500)
                try:
                    s_type = active_app.scholarship.type.value
                except AttributeError:
                    s_type = str(active_app.scholarship.type)

                if active_app.scholarship.results_date:
                    try:
                        delta = now.replace(tzinfo=None) - active_app.scholarship.results_date.replace(tzinfo=None)
                        d_active = max(0, delta.days)
                    except Exception:
                        pass

            career_name = student.career_rel.name if student.career_rel else None

            # 👇 NUEVO: Calculamos cuántas faltas tiene este becario esta semana
            faults = session.exec(
                select(func.count(Attendance.id)).where(
                    Attendance.student_id == student.control_number,
                    Attendance.date >= start_of_week,
                    Attendance.date <= end_of_week,
                    Attendance.status == AttendanceStatus.FALTA
                )
            ).one()

            student_data = StudentReadWithCareer(
                control_number=student.control_number,
                full_name=student.full_name,
                email=student.email,
                career=career_name,
                is_blacklisted=student.is_blacklisted,
                career_rel=CareerRelRead(name=career_name) if career_name else None,
                scholarship_type=s_type,
                days_active=d_active,
                total_services=total_serv,
                released_services=released_serv,
                current_week_faults=faults
            )

            items.append(student_data)

        return PaginatedStudents(total=total, items=items)

    except HTTPException:
        raise
    except Exception as e:
        import traceback
        print("--- ERROR FATAL EN READ_STUDENTS ---")
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=f"Error en BD: {str(e)}")


# --- ALTA MANUAL DE BECARIO ---
@router.post("/manual", response_model=StudentReadWithCareer)
def create_student_manual(
        student_in: StudentManualCreate,
        session: Session = Depends(get_session),
        current_user: User = Depends(get_current_user),
):
    _check_becarios_manager(current_user)

    control = student_in.control_number.strip()
    if not control:
        raise HTTPException(status_code=400, detail="El número de control es obligatorio.")
    if not student_in.full_name.strip():
        raise HTTPException(status_code=400, detail="El nombre completo es obligatorio.")

    if not student_in.career.strip():
        raise HTTPException(status_code=400, detail="La carrera es obligatoria.")

    career_id = _resolve_career_id(session, student_in.career)
    if career_id is None:
        raise HTTPException(status_code=400, detail="La carrera seleccionada no existe en el catálogo.")

    existing = session.get(Student, control)

    if existing:
        # Comportamiento idempotente: si existe, se actualizan los datos
        existing.full_name = student_in.full_name.strip()
        if student_in.email:
            existing.email = student_in.email.strip()
        if student_in.phone_number:
            existing.phone_number = student_in.phone_number.strip()
        existing.career_id = career_id
        session.add(existing)
        session.commit()
        session.refresh(existing)
    else:
        existing = Student(
            control_number=control,
            full_name=student_in.full_name.strip(),
            email=(student_in.email or "").strip(),
            phone_number=student_in.phone_number.strip() if student_in.phone_number else None,
            career_id=career_id,
        )
        session.add(existing)
        session.commit()
        session.refresh(existing)

    career_name = existing.career_rel.name if existing.career_rel else None
    return StudentReadWithCareer(
        control_number=existing.control_number,
        full_name=existing.full_name,
        email=existing.email,
        career=career_name,
        is_blacklisted=existing.is_blacklisted,
        career_rel=CareerRelRead(name=career_name) if career_name else None,
        total_services=0,
        released_services=0,
        current_week_faults=0,
    )


# --- IMPORTACIÓN MASIVA DESDE EXCEL ---
@router.post("/import-excel", response_model=ImportExcelResponse)
def import_students_excel(
        file: UploadFile = File(...),
        session: Session = Depends(get_session),
        current_user: User = Depends(get_current_user),
):
    _check_becarios_manager(current_user)

    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="El archivo debe ser un Excel (.xlsx).")

    try:
        wb = load_workbook(BytesIO(file.file.read()), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="No se pudo leer el archivo Excel. Verifica que sea válido.")

    def normalize(value) -> str:
        if value is None:
            return ""
        text = str(value).strip().lower()
        for accented, plain in (("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"), ("ü", "u")):
            text = text.replace(accented, plain)
        return text

    def clean_header(value) -> str:
        text = normalize(value)
        for ch in (".", ":", "_", "#", "-", "/", "\n", "\t"):
            text = text.replace(ch, " ")
        return " ".join(text.split())

    def find_column(headers: List[str], *aliases) -> Optional[int]:
        for i, h in enumerate(headers):
            if any(alias in h for alias in aliases):
                return i
        return None

    worksheets = wb.worksheets
    if not worksheets:
        raise HTTPException(status_code=400, detail="El archivo está vacío.")

    creados = 0
    actualizados = 0
    duplicados = 0
    servicios = 0
    errores: List[str] = []
    seen: set = set()
    now = datetime.utcnow()
    processed_sheets = 0

    for ws in worksheets:
        sheet_rows = list(ws.iter_rows(values_only=True))
        if not sheet_rows:
            continue

        sheet_label = f"Hoja '{ws.title}'"

        # Buscar la fila de encabezado en cualquier posición (ignora títulos/filas vacías)
        header_index: Optional[int] = None
        col_control = col_name = col_career = col_actividad = col_folio = None

        for idx, candidate in enumerate(sheet_rows):
            cleaned = [clean_header(c) for c in candidate]
            c_control = find_column(cleaned, "no control", "nocontrol", "numero de control", "num de control",
                                    "num control", "no de control", "control", "ctrl", "matricula")
            c_name = find_column(cleaned, "nombre", "alumno", "estudiante", "name")
            if c_control is None or c_name is None:
                continue
            header_index = idx
            col_control = c_control
            col_name = c_name
            col_career = find_column(cleaned, "carrera", "career", "programa", "licenciatura")
            col_actividad = find_column(cleaned, "actividad", "servicio")
            col_folio = find_column(cleaned, "folio")
            break

        if header_index is None:
            errores.append(f"{sheet_label}: se omitió porque no se encontró el encabezado (número de control y nombre).")
            continue

        if col_career is None:
            errores.append(f"{sheet_label}: se omitió porque no tiene la columna de carrera.")
            continue

        processed_sheets += 1

        for row_number, row in enumerate(sheet_rows[header_index + 1:], start=header_index + 2):
            def cell(idx: Optional[int]) -> str:
                if idx is None or idx >= len(row):
                    return ""
                value = row[idx]
                if value is None:
                    return ""
                if isinstance(value, float) and value.is_integer():
                    return str(int(value))
                return str(value).strip()

            control = cell(col_control)
            name = cell(col_name)
            career_name = cell(col_career)
            actividad = cell(col_actividad)
            folio = cell(col_folio)

            if not control and not name:
                continue  # Fila vacía

            if not control:
                errores.append(f"{sheet_label} fila {row_number}: falta el número de control.")
                continue
            if not name:
                errores.append(f"{sheet_label} fila {row_number}: falta el nombre del becario.")
                continue
            if not career_name:
                errores.append(f"{sheet_label} fila {row_number}: falta la carrera del becario.")
                continue

            career_id = _resolve_career_id(session, career_name)
            if career_id is None:
                errores.append(f"{sheet_label} fila {row_number}: la carrera '{career_name}' no existe en el catálogo.")
                continue

            student = _find_student(session, control)
            service_control = student.control_number if student else control

            if control.lower() in seen:
                duplicados += 1
            else:
                seen.add(control.lower())
                if student:
                    updated = False
                    if student.full_name != name:
                        student.full_name = name
                        updated = True
                    if student.career_id != career_id:
                        student.career_id = career_id
                        updated = True
                    if updated:
                        session.add(student)
                    actualizados += 1
                else:
                    session.add(Student(
                        control_number=control,
                        full_name=name,
                        email="",
                        career_id=career_id,
                    ))
                    creados += 1

            # Servicio (opcional): solo se registra si vienen actividad y folio
            if actividad and folio:
                existing_srv = session.exec(
                    select(ServicioBecario).where(ServicioBecario.folio == folio)
                ).first()
                if existing_srv:
                    if _normalize_text(existing_srv.control_number) == _normalize_text(service_control):
                        if existing_srv.actividad != actividad:
                            existing_srv.actividad = actividad
                            session.add(existing_srv)
                        servicios += 1
                    else:
                        errores.append(f"{sheet_label} fila {row_number}: el folio {folio} ya está asignado a otro becario.")
                else:
                    periodo, anio = _derive_periodo_anio(folio, now)
                    session.add(ServicioBecario(
                        control_number=service_control,
                        actividad=actividad,
                        periodo=periodo,
                        anio=anio,
                        estado=ServicioEstado.LIBERADO,
                        folio=folio,
                        liberado_at=now,
                    ))
                    servicios += 1

    if processed_sheets == 0:
        raise HTTPException(
            status_code=400,
            detail="Ninguna hoja del Excel tiene las columnas requeridas: No. de Control, Nombre y Carrera.",
        )

    try:
        session.commit()
    except Exception as e:
        session.rollback()
        raise HTTPException(status_code=500, detail=f"Error guardando los registros: {str(e)}")

    return ImportExcelResponse(
        total=creados + actualizados + duplicados,
        creados=creados,
        actualizados=actualizados,
        duplicados=duplicados,
        servicios=servicios,
        errores=errores,
    )


# --- SERVICIOS DEL BECARIO ---
@router.get("/{control_number}/servicios", response_model=List[ServicioRead])
def get_student_servicios(
        control_number: str,
        session: Session = Depends(get_session),
        current_user: User = Depends(get_current_user),
):
    _check_becarios_manager(current_user)

    servicios = session.exec(
        select(ServicioBecario)
        .where(ServicioBecario.control_number == control_number)
        .order_by(ServicioBecario.created_at.desc(), ServicioBecario.id.desc())
    ).all()

    return [_to_servicio_read(s) for s in servicios]


@router.post("/{control_number}/servicios", response_model=ServicioRead)
def create_student_servicio(
        control_number: str,
        servicio_in: ServicioCreate,
        session: Session = Depends(get_session),
        current_user: User = Depends(get_current_user),
):
    _check_becarios_manager(current_user)
    if not control_number.strip():
        raise HTTPException(status_code=400, detail="Número de control inválido.")

    student = session.get(Student, control_number)
    if not student:
        raise HTTPException(status_code=404, detail="Estudiante no encontrado.")

    if not servicio_in.actividad.strip():
        raise HTTPException(status_code=400, detail="La actividad es obligatoria.")

    if not (2000 <= servicio_in.anio <= 2100):
        raise HTTPException(status_code=400, detail="El año es inválido.")

    servicio = ServicioBecario(
        control_number=control_number,
        actividad=servicio_in.actividad.strip(),
        periodo=servicio_in.periodo,
        anio=servicio_in.anio,
    )

    if servicio_in.liberar:
        servicio.estado = ServicioEstado.LIBERADO

    session.add(servicio)
    session.commit()
    session.refresh(servicio)

    if servicio.estado == ServicioEstado.LIBERADO:
        servicio.folio = _generate_folio(servicio)
        servicio.liberado_at = datetime.utcnow()
        session.add(servicio)
        session.commit()
        session.refresh(servicio)

    return _to_servicio_read(servicio)


@router.patch("/{control_number}/servicios/{servicio_id}", response_model=ServicioRead)
def update_student_servicio(
        control_number: str,
        servicio_id: int,
        servicio_in: ServicioUpdate,
        session: Session = Depends(get_session),
        current_user: User = Depends(get_current_user),
):
    _check_becarios_manager(current_user)

    servicio = session.get(ServicioBecario, servicio_id)
    if not servicio or servicio.control_number != control_number:
        raise HTTPException(status_code=404, detail="Servicio no encontrado.")

    update_data = servicio_in.model_dump(exclude_unset=True)

    # Cambio de estado vía 'liberar'
    if "liberar" in update_data:
        liberar = update_data.pop("liberar")
        if liberar:
            if servicio.estado != ServicioEstado.LIBERADO:
                servicio.estado = ServicioEstado.LIBERADO
                servicio.liberado_at = datetime.utcnow()
                session.add(servicio)
                session.commit()
                session.refresh(servicio)
                if not servicio.folio:
                    servicio.folio = _generate_folio(servicio)
                    session.add(servicio)
        else:
            # Revertir liberación
            servicio.estado = ServicioEstado.PENDIENTE
            servicio.folio = None
            servicio.liberado_at = None

    for key, value in update_data.items():
        setattr(servicio, key, value)

    session.add(servicio)
    session.commit()
    session.refresh(servicio)
    return _to_servicio_read(servicio)


@router.delete("/{control_number}/servicios/{servicio_id}")
def delete_student_servicio(
        control_number: str,
        servicio_id: int,
        session: Session = Depends(get_session),
        current_user: User = Depends(get_current_user),
):
    _check_becarios_manager(current_user)

    servicio = session.get(ServicioBecario, servicio_id)
    if not servicio or servicio.control_number != control_number:
        raise HTTPException(status_code=404, detail="Servicio no encontrado.")

    session.delete(servicio)
    session.commit()
    return {"ok": True, "message": "Servicio eliminado correctamente"}


@router.get("/{control_number}/history")
def get_student_history(
        control_number: str,
        session: Session = Depends(get_session),
        current_user: User = Depends(get_current_user)
):
    if not is_becarios_manager(current_user):
        raise HTTPException(status_code=403, detail="No autorizado")

    applications = session.exec(
        select(ScholarshipApplication)
        .where(ScholarshipApplication.control_number == control_number)
        .options(selectinload(ScholarshipApplication.scholarship))
        .order_by(ScholarshipApplication.created_at.desc())
    ).all()

    return applications


@router.patch("/{control_number}/toggle-blacklist")
def toggle_blacklist(
        control_number: str,
        session: Session = Depends(get_session),
        current_user: User = Depends(get_current_user)
):
    if current_user.role not in [UserRole.ADMIN_SYS, UserRole.ESTRUCTURA] and current_user.area != UserArea.PREVENCION:
        raise HTTPException(status_code=403, detail="No autorizado para vetar alumnos")

    student = session.get(Student, control_number)
    if not student: raise HTTPException(status_code=404, detail="No encontrado")

    student.is_blacklisted = not student.is_blacklisted
    session.add(student)
    session.commit()
    return student